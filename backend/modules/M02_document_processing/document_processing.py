"""PDF/DOCX text extraction; input metadata validation is owned by M1."""
from io import BytesIO
from pathlib import Path
import pymupdf
from docx import Document
from fastapi import HTTPException
from modules.M01_input_sources.input_validation import validate_file_metadata
from modules.M03_nlp_information_extraction.resume_parser import clean_text

def extract_document_text(filename: str | None, content: bytes) -> str:
    extension = validate_file_metadata(filename, content)
    try:
        if extension == ".pdf":
            with pymupdf.open(stream=content, filetype="pdf") as pdf:
                text = "\n".join(page.get_text() for page in pdf)
        else:
            doc = Document(BytesIO(content))
            parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells: parts.append(" | ".join(cells))
            text = "\n".join(parts)
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=422, detail="Invalid or unreadable document.") from exc
    text = clean_text(text)
    if not text.strip():
        raise HTTPException(status_code=422, detail="No readable text found in the document. Scanned PDFs may require OCR.")
    return text
