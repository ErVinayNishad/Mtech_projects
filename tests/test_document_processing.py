from io import BytesIO

import pymupdf
import pytest
from docx import Document
from fastapi import HTTPException

from modules.M02_document_processing.document_processing import extract_document_text


def make_pdf(text: str) -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), text)
    content = doc.tobytes()
    doc.close()
    return content


def make_docx() -> bytes:
    doc = Document()
    doc.add_heading("Technical Skills", level=1)
    doc.add_paragraph("Python and PyTorch Lightning")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Project"
    table.cell(0, 1).text = "Rust API"
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def test_pdf_text_extraction_and_cleaning():
    text = extract_document_text("resume.pdf", make_pdf("Technical Skills\nPython"))
    assert "Technical Skills" in text
    assert "Python" in text


def test_docx_text_extraction_includes_tables():
    text = extract_document_text("resume.docx", make_docx())
    assert "Technical Skills" in text
    assert "PyTorch Lightning" in text
    assert "Project | Rust API" in text


def test_corrupt_pdf_returns_clear_http_error():
    with pytest.raises(HTTPException) as exc:
        extract_document_text("resume.pdf", b"not a PDF")
    assert exc.value.status_code == 422


def test_empty_extracted_document_returns_clear_http_error():
    with pytest.raises(HTTPException) as exc:
        extract_document_text("resume.pdf", make_pdf("   "))
    assert exc.value.status_code == 422
