from io import BytesIO
from pathlib import Path

import pymupdf
from docx import Document
from fastapi import FastAPI, File, Form, HTTPException, UploadFile

from app.services.feedback.candidate_feedback import (
    generate_candidate_feedback,
)
from app.services.resume.resume_parser import (
    clean_text,
    extract_sections,
    extract_personal_details,
)
from app.services.resume.model_skill_extractor import (
    get_claimed_skills_with_gliner,
)
from app.services.job.jd_skill_extractor import extract_required_skills
from app.services.skill_analysis.skill_gap import calculate_skill_gap
from app.services.skill_analysis.initial_analysis import analyze_skill_gap


app = FastAPI(
    title="Resume Skill Verification API",
    description="Resume parsing and skill analysis system",
    version="1.0.0",
)

MAX_FILE_SIZE = 5 * 1024 * 1024


def extract_document_text(filename, content):
    extension = Path(filename).suffix.lower()

    if extension not in {".pdf", ".docx"}:
        raise HTTPException(
            status_code=400,
            detail="Only PDF and DOCX files are supported.",
        )

    if not content:
        raise HTTPException(
            status_code=400,
            detail="File is empty.",
        )

    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(
            status_code=413,
            detail="File size must not exceed 5 MB.",
        )

    try:
        if extension == ".pdf":
            text_parts = []

            with pymupdf.open(
                stream=content,
                filetype="pdf",
            ) as document:
                for page in document:
                    page_text = page.get_text()
                    if page_text:
                        text_parts.append(page_text)

            text = "\n".join(text_parts)

        else:
            document = Document(BytesIO(content))
            text_parts = []

            for paragraph in document.paragraphs:
                value = paragraph.text.strip()
                if value:
                    text_parts.append(value)

            for table in document.tables:
                for row in table.rows:
                    row_text = [
                        cell.text.strip()
                        for cell in row.cells
                        if cell.text.strip()
                    ]

                    if row_text:
                        text_parts.append(" | ".join(row_text))

            text = "\n".join(text_parts)

    except Exception as exc:
        raise HTTPException(
            status_code=422,
            detail="Invalid or unreadable document.",
        ) from exc

    text = clean_text(text)

    if not text.strip():
        raise HTTPException(
            status_code=422,
            detail="No readable text found in the document.",
        )

    return text


def get_job_description_text(job_description, jd_file, jd_content=None):
    has_text = bool(job_description and job_description.strip())
    has_file = jd_file is not None

    if has_text == has_file:
        raise HTTPException(
            status_code=400,
            detail="Provide either JD text or a JD PDF/DOCX file, not both.",
        )

    if has_text:
        return job_description.strip()

    if jd_content is None:
        raise HTTPException(
            status_code=400,
            detail="JD file content was not provided.",
        )

    return extract_document_text(
        jd_file.filename or "",
        jd_content,
    )


def process_resume(filename, content):
    text = extract_document_text(filename, content)
    sections = extract_sections(text)
    personal_details = extract_personal_details(text)

    technical_skills = sections.get("technical_skills", [])

    if isinstance(technical_skills, str):
        technical_skills = [technical_skills]

    skills_text = "\n".join(
        line
        for line in technical_skills
        if isinstance(line, str)
    )

    claimed_skills = get_claimed_skills_with_gliner(
        skills_text,
        section_name="technical_skills",
    )

    return {
        "filename": Path(filename).name,
        "file_type": Path(filename).suffix.lower()[1:],
        "text": text,
        "personal_details": personal_details,
        "sections": sections,
        "claimed_skills": claimed_skills,
    }


@app.get("/")
def home():
    return {"message": "Resume Skill Verification API is running"}


@app.get("/health")
def health_check():
    return {"status": "healthy"}


@app.post("/api/resume/upload")
async def upload_resume(file: UploadFile = File(...)):
    try:
        content = await file.read()
        resume = process_resume(file.filename or "", content)

        return {
            "filename": resume["filename"],
            "file_type": resume["file_type"],
            "text_length": len(resume["text"]),
            "extracted_text": resume["text"],
            "personal_details": resume["personal_details"],
            "sections": resume["sections"],
            "claimed_skills": resume["claimed_skills"],
            "message": "Resume parsed successfully.",
        }
    finally:
        await file.close()


@app.post("/api/job-description/analyze")
async def analyze_job_description(
    job_description: str | None = Form(None),
    jd_file: UploadFile | None = File(None),
):
    try:
        jd_content = await jd_file.read() if jd_file else None

        jd_text = get_job_description_text(
            job_description,
            jd_file,
            jd_content,
        )

        required_skills = extract_required_skills(jd_text)

        return {
            "input_type": (
                "text"
                if job_description and job_description.strip()
                else "file"
            ),
            "required_skills": required_skills,
            "total_required_skills": len(required_skills),
            "job_description_length": len(jd_text),
            "message": "Job description analyzed successfully.",
        }
    finally:
        if jd_file:
            await jd_file.close()


@app.post("/api/skill-gap/analyze")
async def analyze_resume_skill_gap(
    file: UploadFile = File(...),
    job_description: str | None = Form(None),
    jd_file: UploadFile | None = File(None),
):
    try:
        jd_content = await jd_file.read() if jd_file else None

        jd_text = get_job_description_text(
            job_description,
            jd_file,
            jd_content,
        )

        resume_content = await file.read()
        resume = process_resume(file.filename or "", resume_content)

        required_skills = extract_required_skills(jd_text)
        claimed_skills = resume["claimed_skills"]

        gap_result = calculate_skill_gap(
            claimed_skills,
            required_skills,
        )

        initial_result = analyze_skill_gap(gap_result)
        feedback_result = generate_candidate_feedback(initial_result)

        return {
            "filename": resume["filename"],
            "jd_input_type": (
                "text"
                if job_description and job_description.strip()
                else "file"
            ),
            "required_skills": required_skills,
            "claimed_skills": [
                item["skill"] for item in claimed_skills
            ],
            "skill_gap": gap_result,
            "initial_analysis": initial_result,
            "candidate_feedback": feedback_result,
            "message": "Initial skill gap analysis completed.",
        }
    finally:
        await file.close()
        if jd_file:
            await jd_file.close()